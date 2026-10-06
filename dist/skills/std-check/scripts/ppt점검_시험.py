r"""ppt점검.py 가 심어 둔 결함을 모두 잡고, 정상 쪽에서 오탐이 없는지 본다.

ppt점검.py 를 고친 뒤 반드시 돌린다. 결함 8 종(글 넘침 · 잘린 그림 · 밖에 그려진 글 ·
글은 안에 있는 넓은 상자(권고만) · 남은 자리표시 글 · 작은 글자 · 없는 글꼴 · 한글 글꼴
미지정)을 심은 덱을 임시 폴더에 만들어 돌린다.

Rerun: C:\Users\SGJEONG99\anaconda3\envs\pyswmm\python.exe -X utf8 ppt점검_시험.py
"""
import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.join(tempfile.gettempdir(), "ppt점검_시험")   # keep test decks out of the repo
TOOL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ppt점검.py")
DECK = os.path.join(HERE, "시험덱.pptx")
OUT = os.path.join(HERE, "시험덱_결과")
PY = sys.executable

os.makedirs(HERE, exist_ok=True)
if os.path.exists(OUT):
    shutil.rmtree(OUT)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def box(s, x, y, w, h, text, size=20, font="맑은 고딕", ea=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size = Pt(size)
        r.font.name = font
        if ea:
            from pptx.oxml.ns import qn
            rPr = r._r.get_or_add_rPr()
            e = rPr.makeelement(qn("a:ea"), {"typeface": font})
            rPr.append(e)
    return tb


# 1 clean slide with notes -> should raise nothing on this slide
s = prs.slides.add_slide(blank)
box(s, 0.5, 0.4, 12, 1, "강우 사상별 첨두 유출량은 관측 대비 8 % 이내다", 32)
box(s, 0.5, 1.6, 12, 1, "사상 12 건 중 11 건이 기준을 만족한다", 20)
s.notes_slide.notes_text_frame.text = "여기서 11 건을 짚는다"
# 2 overflow text
s = prs.slides.add_slide(blank)
box(s, 0.5, 0.5, 4, 0.5, "넘치는 글 " * 60, 20)
# 3 off the slide: picture cut (error), text drawn past the edge (error),
#   wide box whose text stays inside (advice only, real-deck case 2026-09-22)
s = prs.slides.add_slide(blank)
from PIL import Image
Image.new("RGB", (200, 100), (200, 30, 30)).save(os.path.join(HERE, "red.png"))
s.shapes.add_picture(os.path.join(HERE, "red.png"), Inches(12), Inches(3), Inches(3))
box(s, 12.6, 1.0, 3, 0.6, "밖으로 나간 글", 20)
box(s, 0.5, 5.5, 16, 0.6, "넓은 상자 안의 짧은 글", 20)
# 4 leftover template text + tiny font
s = prs.slides.add_slide(blank)
box(s, 0.5, 0.5, 12, 1, "제목을 입력하십시오", 32)
box(s, 0.5, 2, 12, 1, "출처 표기", 9)
# 5 missing font
s = prs.slides.add_slide(blank)
box(s, 0.5, 0.5, 12, 1, "없는 글꼴", 24, font="Aptos")
# 6 hangul without a:ea
s = prs.slides.add_slide(blank)
box(s, 0.5, 0.5, 12, 1, "동아시아 글꼴 없음", 24, ea=False)
prs.save(DECK)

r = subprocess.run([PY, "-X", "utf8", TOOL, DECK, OUT], capture_output=True,
                   text=True, encoding="utf-8")
print(r.stdout[-3000:], r.stderr[-2000:])
o = r.stdout
expect = {
    "overflow p2": "2쪽  'TextBox 1' 글 넘침",
    "offslide pic p3": "3쪽  'Picture 1' 가 슬라이드 밖으로",
    "text drawn outside p3": "3쪽  'TextBox 2' 글이 슬라이드 밖에 그려진다",
    "wide box advice p3": "[권고] 3쪽  'TextBox 3' 글상자가 슬라이드 밖으로",
    "leftover p4": "4쪽  'TextBox 1' 에 자리표시 글",
    "tiny p4": "4쪽  12pt 미만 글 1 곳 (최소 9pt, 'TextBox 2'",
    "font Aptos": "없는 글꼴 'Aptos'",
    "no ea p6": "한글에 동아시아 글꼴(a:ea) 미지정 1 곳 (6쪽)",
}
miss = [k for k, v in expect.items() if v not in o]
false_pos = [ln for ln in o.splitlines() if ln.startswith("[") and "1쪽" in ln
             and "[결과]" not in ln]
print("missed:", miss)
print("false positives on clean slide:", false_pos)
assert r.returncode == 1, "should fail with errors, got %d" % r.returncode
assert not miss, miss
assert not false_pos, false_pos
assert os.path.exists(os.path.join(OUT, "그림", "00_전체보기.png"))
assert len([f for f in os.listdir(os.path.join(OUT, "그림")) if f.startswith("슬라이드_")]) == 6
print("OK: all 8 planted defects caught, clean slide silent")
