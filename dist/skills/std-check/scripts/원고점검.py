r"""논문 원고 점검. 사람이 교정하기 전에 기계로 거른다.

왜 있나
  원고를 고칠 때마다 본문 인용과 참고문헌 목록이 어긋나고, 그림 번호가 비고, 약어가
  정의 전에 나온다. 심사자는 이것부터 본다. 또 AI 에게 문헌을 찾게 하면 있지도 않은
  논문이나 DOI 를 지어낸다. 눈으로는 못 잡으니 기계로 맞대고, 참고문헌은 Crossref 에
  실제로 있는지 묻는다.

무엇을 보나
  오류 (종료코드 1)
    본문에 인용했는데 목록에 없음 · 번호 인용이 목록 수를 넘음
    그림 · 표 캡션 번호가 비거나 겹침 · 본문이 없는 그림 · 표 번호를 부름
    (--crossref) DOI 가 Crossref 에 없음 · DOI 의 논문 제목 · 연도가 목록과 다름
  권고
    목록에 있는데 본문에서 인용하지 않음 · 캡션은 있는데 본문에서 부르지 않는 그림 · 표
    그림 · 표를 처음 부르는 순서가 번호 순이 아님
    약어가 정의 전에 쓰이거나 정의가 없음 · 숫자와 단위 사이 띄어쓰기 없음(10mm)
    (--crossref) DOI 가 없고 제목으로도 찾지 못함. 손으로 확인한다

실행
  C:\Users\SGJEONG99\anaconda3\envs\pyswmm\python.exe -X utf8 원고점검.py <원고.docx|.md|.txt> [--crossref]
  결과는 원고 옆 <이름>_원고점검.md. --crossref 는 참고문헌마다 api.crossref.org 에
  묻는다(참고문헌 문자열만 보낸다. 원고 본문은 보내지 않는다). 결과는 옆의
  <이름>_crossref.json 에 저장해 다시 돌릴 때 재사용한다.
시험
  원고점검_시험.py 가 심어 둔 결함을 모두 잡는지 본다.
"""
import difflib
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

# ------------------------------------------------------------------ 기준

참고문헌제목 = re.compile(r"^\s*(\d+\.?\s*)?(references|reference|bibliography|literature cited|"
                     r"참고\s*문헌|참고문헌|인용\s*문헌)\s*$", re.I)
# 참고문헌 뒤에 오는 절(부록 · 그림 목록 등)에서 목록을 끊는다
# thesis back matter too: the Korean abstract after References was read as an entry
목록끝 = re.compile(r"^\s*(appendix|appendices|supplementary|figure captions|tables?|list of figures|"
                 r"abstract|acknowledg|국문\s*초록|초\s*록|요\s*약|주제어|감사의\s*글|"
                 r"부록|그림\s*목록|표\s*목록)", re.I)

연도 = r"(?:19|20)\d{2}[a-z]?"
# 영문 저자-연도 인용. 괄호 안 여러 개는 ; 로 나뉜다
괄호인용 = re.compile(r"\(([^()]*?" + 연도 + r"[^()]*?)\)")
서술인용 = re.compile(r"([A-Z][A-Za-z\u00C0-\u017F'\-]+)(?:\s+(?:et\s+al\.|and|&)\s*"
                  r"(?:[A-Z][A-Za-z\u00C0-\u017F'\-]+)?)?\s*\((" + 연도 + r")\)")
한글서술인용 = re.compile(r"([가-힣]{2,4})(?:\s*(?:등|외|과|와)\s*(?:[가-힣]{2,4})?)?\s*\((" + 연도 + r")\)")
번호인용 = re.compile(r"\[(\d+(?:\s*[-,\u2013]\s*\d+)*)\]")

그림말 = r"(?:Fig\.|Figs\.|Figure|Figures|그림)"
표말 = r"(?:Table|Tables|표)"
캡션 = re.compile(r"^\s*(" + 그림말 + r"|" + 표말 + r")\s*(\d+)\s*[.:]")
부름 = re.compile(r"(" + 그림말 + r"|" + 표말 + r")\s*(\d+)(?:\s*[a-z]\b|\([a-z]\))?"
                 r"(?:\s*(?:and|&|,|-|\u2013|및|과|와)\s*(\d+))?")

약어 = re.compile(r"\b([A-Z][A-Z0-9]{1,6}s?)\b")
약어정의 = re.compile(r"\(\s*([A-Z][A-Z0-9]{1,6}s?)\s*[;,)]")
# 정의 없이 써도 되는 것. 단위 · 국가 · 서식 기호
약어예외 = {"USA", "UK", "EU", "DOI", "ID", "PDF", "URL", "CPU", "GPU", "RAM", "OK", "II", "III",
         "IV", "VI", "VII", "VIII", "IX", "XI", "XII", "AM", "PM", "UTC", "KST", "GMT", "CRAN",
         "HTTP", "HTTPS", "XML", "CSV", "API", "PC", "NA", "SI", "MIT", "USD", "KRW"}
# no 1-letter units except m: "1970s" is a decade, not seconds (real manuscript)
단위붙음 = re.compile(r"\b(\d+(?:\.\d+)?)(mm|cm|km|m|ha|km2|m3|mL|mg|kg|min|hr|yr|"
                   r"mm/h|mm/hr|m3/s|m/s|mg/L|°C|kPa|MPa)\b")


# ------------------------------------------------------------------ 읽기

def 문단들(path):
    """Paragraphs and table-cell texts, kept apart. Table cells appended after the
    paragraphs used to land behind the References heading and were read as 165
    bogus reference entries (real JoH manuscript, 2026-09-22)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        칸 = []
        for t in d.tables:                   # captions and citations can live in tables
            for row in t.rows:
                for c in row.cells:
                    칸.extend(p.text for p in c.paragraphs)
        return [p.text for p in d.paragraphs], 칸
    return open(path, encoding="utf-8", errors="ignore").read().splitlines(), []


def 나누기(ps):
    """Split paragraphs into body and reference entries."""
    k = next((i for i, p in enumerate(ps) if 참고문헌제목.match(p)), None)
    if k is None:
        return ps, []
    목록 = []
    for p in ps[k + 1:]:
        if 목록끝.match(p):
            break
        s = p.strip()
        if not s:
            continue
        # a reference broken over two paragraphs: glue a fragment that does not start
        # like a new entry (author surname or [n]) onto the previous one
        # multi-word surnames count as a new entry: "Kheshti Azar, M." was glued onto
        # the previous reference in the real manuscript (2026-09-22)
        if 목록 and not re.match(r"^(\[\d+\]|\d+\.|[A-Z][A-Za-zÀ-ſ'\-]*(?:\s+[A-Z][A-Za-zÀ-ſ'\-]*)*,|"
                                 r"(?:van|von|de|der|den|di|da|du|la|le)\s|[가-힣]{2,4}[,.\s])", s):
            목록[-1] += " " + s
            continue
        목록.append(s)
    return ps[:k], 목록


# ------------------------------------------------------------------ 인용

def 열쇠(성, 해):
    return (성.lower().replace("\u2019", "'"), 해.lower())


def 본문인용(본문):
    """(surname, year) keys cited in the body, plus numbered citations."""
    키, 번호 = {}, set()
    for p in 본문:
        for m in 괄호인용.finditer(p):
            for 조각 in re.split(r";", m.group(1)):
                y = re.search(연도, 조각)
                a = re.search(r"([A-Z][A-Za-z\u00C0-\u017F'\-]+|[가-힣]{2,4})", 조각)
                if y and a and not re.match(r"^(Fig|Table|Eq|Figure|See|e\.g|i\.e|cf)", a.group(1)):
                    # (Smith et al., 2009a, b) -> both years
                    for yy in re.findall(연도 + r"|(?<=[,\s])[a-z](?=[,;\s)]|$)", 조각):
                        if len(yy) == 1:
                            yy = y.group(0)[:4] + yy
                        키.setdefault(열쇠(a.group(1), yy), p[:60])
        for m in 서술인용.finditer(p):
            키.setdefault(열쇠(m.group(1), m.group(2)), p[:60])
        for m in 한글서술인용.finditer(p):
            키.setdefault(열쇠(m.group(1), m.group(2)), p[:60])
        for m in 번호인용.finditer(p):
            for 덩이 in re.split(r"\s*,\s*", m.group(1)):
                ab = re.split(r"\s*[-\u2013]\s*", 덩이)
                if len(ab) == 2 and ab[0].isdigit() and ab[1].isdigit():
                    번호.update(range(int(ab[0]), int(ab[1]) + 1))
                elif ab[0].isdigit():
                    번호.add(int(ab[0]))
    return 키, 번호


def 목록열쇠(항목):
    """(set of surname words of the first author, year) from a reference entry.
    The first author is everything before the first comma, so particles and compound
    surnames survive: 'van Bijnen, M.' -> {van, bijnen}, 'Kheshti Azar, M.' -> {kheshti, azar}."""
    s = re.sub(r"^\s*(\[\d+\]|\d+\.)\s*", "", 항목)
    # "(submitted)", "(in press)" stand in for a year in theses and revisions
    y = re.search(r"\(?(" + 연도 + r"|submitted|in press|in review|under review|forthcoming)\)?", s, re.I)
    a = re.match(r"\s*([^,(.]+?)\s*[,(.]", s)
    if not (a and y):
        return None
    말 = {w.lower().replace("’", "'")
         for w in re.findall(r"[A-Za-zÀ-ſ'\-]+|[가-힣]+", a.group(1))}
    return (frozenset(말), y.group(1).lower()) if 말 else None


def 맞나(인용, 항목키):
    return 항목키 is not None and 인용[0] in 항목키[0] and 인용[1][:4] == 항목키[1][:4]


# ------------------------------------------------------------------ 그림 · 표 · 약어 · 단위

def 번호점검(본문, 오류, 권고):
    캡션번호 = {"그림": [], "표": []}
    부른번호 = {"그림": [], "표": []}
    for p in 본문:
        m = 캡션.match(p)
        if m:
            종류 = "표" if re.match(표말, m.group(1)) else "그림"
            캡션번호[종류].append(int(m.group(2)))
            continue
        for m in 부름.finditer(p):
            종류 = "표" if re.match(표말, m.group(1)) else "그림"
            a = int(m.group(2))
            부른번호[종류].append(a)
            if m.group(3):
                b = int(m.group(3))
                if "-" in m.group(0) or "\u2013" in m.group(0):
                    부른번호[종류].extend(range(a + 1, b + 1))
                else:
                    부른번호[종류].append(b)
    for 종류 in ("그림", "표"):
        c = 캡션번호[종류]
        if not c:
            continue
        겹침 = sorted({n for n in c if c.count(n) > 1})
        if 겹침:
            오류.append("%s 캡션 번호가 겹친다: %s" % (종류, 겹침))
        빔 = sorted(set(range(1, max(c) + 1)) - set(c))
        if 빔:
            오류.append("%s 캡션 번호가 빈다: %s (있는 것 %s)" % (종류, 빔, sorted(set(c))))
        없는것 = sorted(set(부른번호[종류]) - set(c))
        if 없는것:
            오류.append("본문이 없는 %s 을 부른다: %s" % (종류, 없는것))
        안부름 = sorted(set(c) - set(부른번호[종류]))
        if 안부름:
            권고.append("캡션은 있는데 본문에서 부르지 않는 %s: %s" % (종류, 안부름))
        처음 = []
        for n in 부른번호[종류]:
            if n not in 처음:
                처음.append(n)
        if 처음 != sorted(처음):
            권고.append("%s 을 처음 부르는 순서가 번호 순이 아니다: %s" % (종류, 처음[:12]))
    return 캡션번호


def 약어점검(본문, 권고):
    정의된곳, 처음쓴곳 = {}, {}
    for i, p in enumerate(본문):
        if 캡션.match(p):
            continue
        글자 = re.sub(r"[^A-Za-z]", "", p)
        if 글자 and 글자.isupper() and len(p) < 80:     # "DATA AVAILABILITY" style heading
            continue
        for m in 약어정의.finditer(p):
            정의된곳.setdefault(m.group(1).rstrip("s"), i)
        for m in 약어.finditer(p):
            a = m.group(1).rstrip("s") if len(m.group(1)) > 2 else m.group(1)
            if a in 약어예외 or re.search(r"\d", a):     # ids like MH0126, PC1
                continue
            처음쓴곳.setdefault(a, (i, p[max(0, m.start() - 20):m.end() + 20]))
    늦음 = sorted(a for a, (i, _) in 처음쓴곳.items() if a in 정의된곳 and 정의된곳[a] > i)
    없음 = sorted(a for a in 처음쓴곳 if a not in 정의된곳)
    if 늦음:
        권고.append("정의보다 먼저 쓰인 약어 %d 개: %s" % (len(늦음), ", ".join(늦음[:15])))
    if 없음:
        권고.append("정의(괄호 풀이)를 찾지 못한 약어 %d 개: %s. 널리 아는 것이면 무시한다"
                  % (len(없음), ", ".join(없음[:20])))


def 단위점검(본문, 권고):
    걸린것 = []
    for p in 본문:
        걸린것 += [m.group(0) for m in 단위붙음.finditer(p)]
    if 걸린것:
        권고.append("숫자와 단위 사이 띄어쓰기 없음 %d 곳(SI 는 띄운다): %s"
                  % (len(걸린것), ", ".join(sorted(set(걸린것))[:12])))


# ------------------------------------------------------------------ Crossref

def 가져오기(url, 머리):
    """GET JSON; on HTTP 429 wait (Retry-After or growing pause) and retry.
    The real 48-reference manuscript hit 429 at 0.2 s spacing (2026-09-22)."""
    for k in range(5):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=머리), timeout=30) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            if e.code != 429 or k == 4:
                raise
            쉼 = e.headers.get("Retry-After")
            time.sleep(float(쉼) if 쉼 and 쉼.replace(".", "").isdigit() else 2.0 * (k + 1))


def 크로스레프(항목, 캐시):
    if 항목 in 캐시:
        return 캐시[항목]
    doi = re.search(r"10\.\d{4,9}/[^\s,;]+", 항목)
    머리 = {"User-Agent": "claude_play-manuscript-check/0.1"}   # no email: not ours to send
    try:
        if doi:
            d = doi.group(0).rstrip(".")
            url = "https://api.crossref.org/works/" + urllib.parse.quote(d)
            try:
                m = 가져오기(url, 머리)["message"]
                결과 = {"방식": "doi", "doi": d, "찾음": True}
            except urllib.error.HTTPError as e:
                if e.code == 404:
                    결과 = {"방식": "doi", "doi": d, "찾음": False}
                    캐시[항목] = 결과
                    return 결과
                raise
        else:
            q = urllib.parse.urlencode({"query.bibliographic": 항목[:300], "rows": 1})
            url = "https://api.crossref.org/works?" + q
            it = 가져오기(url, 머리)["message"]["items"]
            if not it:
                결과 = {"방식": "검색", "찾음": False}
                캐시[항목] = 결과
                return 결과
            m = it[0]
            결과 = {"방식": "검색", "찾음": True, "doi": m.get("DOI")}
        제목 = (m.get("title") or [""])[0]
        # online-first and print years often differ by one (Zhang 2009 print, 2008 online)
        해들 = []
        for k in ("issued", "published-print", "published-online", "published"):
            dp = (m.get(k) or {}).get("date-parts")
            if dp and dp[0] and dp[0][0] and dp[0][0] not in 해들:
                해들.append(dp[0][0])
        결과.update({"제목": 제목, "연도": 해들[0] if 해들 else None, "연도들": 해들,
                   "첫저자": ((m.get("author") or [{}])[0]).get("family")})
    except Exception as ex:
        결과 = {"방식": "오류", "찾음": None, "오류": "%s: %s" % (type(ex).__name__, ex)}
    캐시[항목] = 결과
    # Crossref PUBLIC pool since 2025-12-01 (we send no mailto, so not the polite pool):
    # single lookups 5/s, list queries 1/s, concurrency 1 (crossref.org blog 2025-11-05)
    time.sleep(1.05 if 결과.get("방식") == "검색" else 0.25)
    return 결과


def 닮음(a, b):
    f = lambda s: re.sub(r"[^a-z0-9가-힣 ]", "", (s or "").lower())
    a, b = f(a), f(b)
    if not a or not b:
        return 0.0
    if b in a:                                   # entry contains the whole title
        return 1.0
    return difflib.SequenceMatcher(None, a, b).ratio()


def 참고문헌대조(목록, 캐시경로, 오류, 권고):
    캐시 = json.load(open(캐시경로, encoding="utf-8")) if os.path.exists(캐시경로) else {}
    확인, 물음 = 0, 0
    for k, 항목 in enumerate(목록, 1):
        r = 크로스레프(항목, 캐시)
        물음 += 1
        짧게 = 항목[:70]
        if r["찾음"] is None:
            권고.append("[%d] Crossref 조회 실패(%s): %s" % (k, r.get("오류"), 짧게))
            continue
        if r["방식"] == "doi" and not r["찾음"]:
            오류.append("[%d] DOI 가 Crossref 에 없다(지어낸 DOI 일 수 있다) %s: %s" % (k, r["doi"], 짧게))
            continue
        if not r["찾음"]:
            권고.append("[%d] 찾지 못함. 손으로 확인한다: %s" % (k, 짧게))
            continue
        s = 닮음(항목, r.get("제목"))
        해들 = r.get("연도들") or ([r["연도"]] if r.get("연도") else [])
        해맞음 = not 해들 or any(str(h) in 항목 for h in 해들)
        if r["방식"] == "doi":
            if s < 0.6 or not 해맞음:
                오류.append("[%d] DOI 의 논문이 목록과 다르다 (제목 닮음 %.2f, Crossref 연도 %s): %s | Crossref: %s"
                          % (k, s, r.get("연도"), 짧게, (r.get("제목") or "")[:60]))
            else:
                확인 += 1
        else:
            if s >= 0.8 and 해맞음:
                확인 += 1
                권고.append("[%d] DOI 가 없다. Crossref 에서 찾은 DOI: %s" % (k, r.get("doi")))
            elif s >= 0.8:
                권고.append("[%d] 제목은 맞는데 연도가 다르다(Crossref %s). 판 · 쇄를 확인한다: %s"
                          % (k, 해들, 짧게))
            else:
                권고.append("[%d] 제목 검색으로 확실한 짝을 못 찾음(닮음 %.2f). 손으로 확인한다: %s"
                          % (k, s, 짧게))
    json.dump(캐시, open(캐시경로, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return 확인, 물음


# ------------------------------------------------------------------ 본체

def main(path, crossref=False):
    assert os.path.isfile(path), "파일이 없다: %s" % path
    path = os.path.abspath(path)
    ps, 칸 = 문단들(path)
    본문, 목록 = 나누기(ps)
    assert 본문, "본문이 0 문단이다"
    본문 = 본문 + 칸                             # captions and citations inside tables
    오류, 권고 = [], []
    if not 목록:
        권고.append("참고문헌 절을 찾지 못했다(References · 참고문헌 제목 문단). 인용 대조를 건너뛴다")

    키, 번호 = 본문인용(본문)
    목록키 = [목록열쇠(x) for x in 목록]
    번호식 = bool(번호) and len(번호) >= len(키)
    if 목록 and 번호식:
        넘침 = sorted(n for n in 번호 if n > len(목록))
        if 넘침:
            오류.append("번호 인용이 목록 수(%d)를 넘는다: %s" % (len(목록), 넘침))
        안씀 = sorted(set(range(1, len(목록) + 1)) - 번호)
        if 안씀:
            권고.append("본문에서 번호로 부르지 않는 참고문헌: %s" % 안씀)
    elif 목록:
        # match on year prefix too: (Smith, 2009) vs list "Smith, 2009a"
        for k, 어디 in sorted(키.items()):
            if not any(맞나(k, 항목키) for 항목키 in 목록키):
                오류.append("본문 인용이 목록에 없다: %s (%s) | %s" % (k[0], k[1], 어디))
        for x, 항목키 in zip(목록, 목록키):
            if 항목키 and not any(맞나(c, 항목키) for c in 키):
                권고.append("목록에 있는데 본문에서 인용하지 않았다: %s" % x[:70])
        못읽음 = [x for x, k in zip(목록, 목록키) if not k]
        if 못읽음:
            권고.append("첫 저자 · 연도를 읽지 못한 목록 항목 %d 개: %s" % (len(못읽음), 못읽음[0][:60]))

    캡션번호 = 번호점검(본문, 오류, 권고)
    약어점검(본문, 권고)
    단위점검(본문, 권고)

    cr = ""
    if crossref and 목록:
        확인, 물음 = 참고문헌대조(목록, os.path.splitext(path)[0] + "_crossref.json", 오류, 권고)
        cr = " · Crossref 확인 %d/%d" % (확인, 물음)

    요약 = ("본문 %d 문단 · 참고문헌 %d 개 · 본문 인용 %s · 그림 캡션 %d · 표 캡션 %d%s"
          % (len(본문), len(목록), ("번호 %d 개" % len(번호)) if 번호식 else ("%d 건" % len(키)),
             len(캡션번호["그림"]), len(캡션번호["표"]), cr))
    줄 = ["# 원고 점검", "", "대상 `%s`" % path, "", 요약,
          "오류 %d 건 · 권고 %d 건" % (len(오류), len(권고)), "", "## 오류", ""] + \
         ["- " + x for x in 오류 or ["없음"]] + ["", "## 권고", ""] + ["- " + x for x in 권고 or ["없음"]] + \
         ["", "## 기계가 못 보는 것", "",
          "인용한 문헌이 정말 그 말을 했는지, 수치가 원장과 같은지, 주장이 결과를 넘지 않는지는",
          "사람이 원문을 열어 본다. `13_논문.md` 의 점검 절.", ""]
    out = os.path.splitext(path)[0] + "_원고점검.md"
    open(out, "w", encoding="utf-8").write("\n".join(줄))
    print(요약)
    for x in 오류:
        print("[오류]", x)
    for x in 권고:
        print("[권고]", x)
    print("결과", out)
    print("[결과] 오류 %d 건 · 권고 %d 건" % (len(오류), len(권고)))
    return 1 if 오류 else 0


if __name__ == "__main__":
    인자 = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(인자) != 1:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(인자[0], crossref="--crossref" in sys.argv))
