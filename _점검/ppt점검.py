r"""발표자료(pptx) 점검. 눈으로 넘기기 전에 기계로 걸러낸다.

왜 있나
  AI 로 만든 슬라이드는 글이 상자를 넘치거나 슬라이드 밖으로 나가도, 양식의
  "제목을 입력하십시오" 가 남아도 파일은 멀쩡히 열린다. python-pptx 로 저장만
  하고 끝내면 아무도 모른다. 2026-09-22 실측: 넘친 제목이 슬라이드 위로 빠져
  본문과 겹쳤는데 python-pptx 는 아무 말도 하지 않았다.
  그래서 실제 PowerPoint(COM)로 열어 그려진 크기를 재고, 쪽마다 PNG 로 내보내
  사람이 전체보기 한 장으로 훑게 한다. LibreOffice 렌더링과 달리 글꼴 대체가
  없으므로 넘침 판정을 믿을 수 있다.

무엇을 보나
  오류 (종료코드 1)   글 넘침 · 슬라이드 밖 도형 · 남은 자리표시 글 · 없는 글꼴
  권고               작은 글자 · 자동 맞춤으로 줄어든 글 · 한글 글꼴(a:ea) 미지정
                     · 그림 한 장뿐인 쪽 · 빈 자리표시자 · 16:9 아님
                     · 슬라이드당 글자 수 과다 · 발표자 노트 없음
  시험               심어 둔 결함 6 종을 모두 잡고 정상 쪽은 조용한지 2026-09-22 확인

실행
  C:\Users\SGJEONG99\anaconda3\envs\pyswmm\python.exe -X utf8 ppt점검.py <발표.pptx> [출력폴더]
  출력폴더를 안 주면 pptx 옆에 <이름>_점검\ 을 만든다. 덮어쓰지 않고 _v2 로 올린다.
  --그림없이   PowerPoint 를 부르지 않는다(넘침 · PNG 판정 생략)
"""
import os
import re
import sys
import warnings

warnings.filterwarnings("error", category=UserWarning, module="pptx")

# ------------------------------------------------------------------ 기준

최소글자pt = 12          # 이보다 작으면 권고. 캡션 10-12, 본문 14 이상이 보통이다
슬라이드당글자 = 350     # 공백 제외. 넘으면 읽는 슬라이드가 된다
넘침여유pt = 2           # 렌더링 반올림
PNG폭 = 1600

# 양식에 남기 쉬운 글. 영어 · 한국어 PowerPoint 기본 문구와 흔한 채움말
남은글 = re.compile(
    r"lorem|ipsum|\bx{3,}\b|\bTODO\b|\bTBD\b|\[insert|click to add|"
    r"제목을 입력|텍스트를 입력|내용을 입력|클릭하여|여기에 입력|부제목 입력|"
    r"\?\?\?|○○|OOO", re.I)

# PowerPoint 는 파일에 적힌 이름을 돌려주고, 설치 목록(tkinter)은 한글 이름을 준다
글꼴별칭 = {
    "Malgun Gothic": "맑은 고딕", "Gulim": "굴림", "Dotum": "돋움",
    "Batang": "바탕", "Gungsuh": "궁서", "NanumGothic": "나눔고딕",
    "NanumMyeongjo": "나눔명조", "NanumBarunGothic": "나눔바른고딕",
    "NanumSquare": "나눔스퀘어", "HCR Dotum": "함초롬돋움", "HCR Batang": "함초롬바탕",
}
테마글꼴 = re.compile(r"^\+(mj|mn)-")    # +mn-lt 같은 테마 참조. 실제 글꼴이 아니다


# ------------------------------------------------------------------ 준비

def 출력폴더(pptx, 지정):
    if 지정:
        os.makedirs(지정, exist_ok=True)
        return 지정
    기본 = os.path.splitext(pptx)[0] + "_점검"
    d, k = 기본, 1
    while os.path.exists(d):                 # 덮어쓰지 않는다
        k += 1
        d = "%s_v%d" % (기본, k)
    os.makedirs(d)
    return d


def 설치글꼴():
    import tkinter
    import tkinter.font as tf
    r = tkinter.Tk()
    r.withdraw()
    f = {x.lstrip("@") for x in tf.families()}
    r.destroy()
    assert len(f) > 20, "글꼴 목록이 비었다: %d" % len(f)
    return f


def 쪽목록(쪽):
    """[1,2,3,5,7,8] -> '1-3, 5, 7-8'"""
    묶음, 시작 = [], None
    for k, n in enumerate(쪽):
        if 시작 is None:
            시작 = n
        if k + 1 == len(쪽) or 쪽[k + 1] != n + 1:
            묶음.append(str(시작) if 시작 == n else "%d-%d" % (시작, n))
            시작 = None
    return ", ".join(묶음)


# ------------------------------------------------------------------ 1. 파일만 보고

def 파일점검(pptx):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    한글 = re.compile(r"[가-힣]")

    prs = Presentation(pptx)
    W, H = prs.slide_width, prs.slide_height
    오류, 권고, 글꼴 = [], [], set()
    비 = W / H
    if abs(비 - 16 / 9) > 0.01:
        권고.append("전체  화면비 %.3f (16:9 는 1.778). %.2f x %.2f in"
                  % (비, Emu(W).inches, Emu(H).inches))

    쪽수 = 0
    노트없음, ea없음 = [], []
    for i, s in enumerate(prs.slides, 1):
        쪽수 += 1
        글자 = 0
        작은글 = []
        for sh in s.shapes:
            # 슬라이드 밖. 자리표시자는 위치가 비어 있을 수 있어 None 을 거른다
            if None not in (sh.left, sh.top, sh.width, sh.height):
                # overshoot in pt on each side; a hair past the edge is not a defect
                밖 = max(-sh.left, -sh.top, sh.left + sh.width - W, sh.top + sh.height - H) / 12700
                if 밖 > 넘침여유pt and sh.has_text_frame:
                    # A wide text box whose text sits inside is not visibly wrong (real deck,
                    # 2026-09-22: caption box 132 pt past the edge, text fully visible).
                    # The rendered-text check in 그려점검 decides; here only advise
                    권고.append("%d쪽  '%s' 글상자가 슬라이드 밖으로 %.0f pt 걸친다. 글은 안에 있는지 PNG 로 본다"
                              % (i, sh.name, 밖))
                elif 밖 > 넘침여유pt:
                    오류.append("%d쪽  '%s' 가 슬라이드 밖으로 %.0f pt 나가 잘린다" % (i, sh.name, 밖))
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text
            글자 += len(re.sub(r"\s", "", t))
            m = 남은글.search(t)
            if m:
                오류.append("%d쪽  '%s' 에 자리표시 글이 남았다: %r" % (i, sh.name, m.group(0)))
            # Autofit that PowerPoint actually applied leaves fontScale in the file.
            # A read-only COM open does not shrink (measured 2026-09-22), so an unsaved
            # python-pptx deck shows up as overflow above instead
            na = sh.text_frame._txBody.bodyPr.find(qn("a:normAutofit"))
            if na is not None and na.get("fontScale"):
                권고.append("%d쪽  '%s' 자동 맞춤으로 글자가 %.0f %% 로 줄었다. 글을 줄이거나 나눈다"
                          % (i, sh.name, int(na.get("fontScale")) / 1000))
            if sh.is_placeholder and not t.strip():
                권고.append("%d쪽  빈 자리표시자 '%s' (발표 화면에는 안 보이지만 지운다)"
                          % (i, sh.name))
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        글꼴.add(r.font.name)
                        # font.name only writes <a:latin>. Hangul follows the theme's
                        # East Asian font (+mn-ea) unless <a:ea> is set (measured 2026-09-22)
                        rPr = r._r.rPr
                        if 한글.search(r.text) and (rPr is None or rPr.find(qn("a:ea")) is None):
                            ea없음.append((i, sh.name, r.font.name))
                    if r.font.size and r.font.size.pt < 최소글자pt and r.text.strip():
                        작은글.append((r.font.size.pt, sh.name, r.text.strip()[:15]))
        if 작은글:                              # one line per slide, not per run
            작은글.sort()
            권고.append("%d쪽  %dpt 미만 글 %d 곳 (최소 %.0fpt, '%s' 의 '%s')"
                      % (i, 최소글자pt, len(작은글), 작은글[0][0], 작은글[0][1], 작은글[0][2]))
        # Image-only slide: typical of AI tools that render a whole slide as one
        # picture (NotebookLM, Marp/Slidev default pptx). Text cannot be edited
        그림수 = sum(1 for sh in s.shapes if sh.shape_type == 13)   # MSO_SHAPE_TYPE.PICTURE
        if 글자 == 0 and 그림수:
            권고.append("%d쪽  글자 0 자 · 그림 %d 개. 쪽 전체가 그림이면 고칠 수 없다" % (i, 그림수))
        if 글자 > 슬라이드당글자:
            권고.append("%d쪽  글자 %d 자 (기준 %d). 읽는 슬라이드가 된다. 나누거나 노트로 옮긴다"
                      % (i, 글자, 슬라이드당글자))
        if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
            노트없음.append(i)
    assert 쪽수 > 0, "슬라이드가 0 장이다"
    # Deck-wide advisories go on one line each; per-slide repetition buried the real
    # findings on a 25-slide deck (2026-09-22: 25 + 30 lines of the same thing)
    if 노트없음:
        권고.append("전체  발표자 노트 없는 쪽 %d/%d: %s" % (len(노트없음), 쪽수, 쪽목록(노트없음)))
    if ea없음:
        쪽 = sorted({x[0] for x in ea없음})
        권고.append("전체  한글에 동아시아 글꼴(a:ea) 미지정 %d 곳 (%s쪽). 지정한 글꼴(%s)은 영문에만 "
                  "걸리고 한글은 테마 글꼴로 나온다. 의도한 것이 아니면 a:ea 를 넣는다"
                  % (len(ea없음), 쪽목록(쪽), " · ".join(sorted({x[2] for x in ea없음}))))
    return 쪽수, 오류, 권고, 글꼴


# ------------------------------------------------------------------ 2. PowerPoint 로 그려 보고

def 도형들(shapes):
    for sh in shapes:
        if sh.Type == 6:                     # msoGroup
            yield from 도형들(sh.GroupItems)
        else:
            yield sh


def 그려점검(pptx, 그림):
    import win32com.client as win32
    오류, 권고, 글꼴, png = [], [], set(), []
    app = win32.Dispatch("PowerPoint.Application")
    d = app.Presentations.Open(os.path.abspath(pptx), ReadOnly=True,
                               Untitled=False, WithWindow=False)
    try:
        W, H = d.PageSetup.SlideWidth, d.PageSetup.SlideHeight
        for i, sl in enumerate(d.Slides, 1):
            for sh in 도형들(sl.Shapes):
                if not sh.HasTextFrame or not sh.TextFrame2.HasText:
                    continue
                tf = sh.TextFrame2
                tr = tf.TextRange
                # 그려진 글이 상자보다 크면 넘침. 위로 빠지는 경우도 있어 위치도 본다
                if tr.BoundHeight > sh.Height + 넘침여유pt:
                    오류.append("%d쪽  '%s' 글 넘침: 글 높이 %.0f > 상자 %.0f pt"
                              % (i, sh.Name, tr.BoundHeight, sh.Height))
                x0, y0 = tr.BoundLeft, tr.BoundTop
                x1, y1 = x0 + tr.BoundWidth, y0 + tr.BoundHeight
                if x1 <= 0 or y1 <= 0 or x0 >= W or y0 >= H:
                    # entirely off the canvas: invisible leftover, not a visible defect
                    # (real deck 2026-09-22: a parked text box below the slide)
                    권고.append("%d쪽  '%s' 가 통째로 슬라이드 밖에 있어 안 보인다. 남은 개체면 지운다: %s"
                              % (i, sh.Name, tr.Text[:20]))
                elif (y0 < -넘침여유pt or y1 > H + 넘침여유pt
                        or x0 < -넘침여유pt or x1 > W + 넘침여유pt):
                    오류.append("%d쪽  '%s' 글이 슬라이드 밖에 그려진다" % (i, sh.Name))
            p = os.path.join(그림, "슬라이드_%02d.png" % i)
            sl.Export(p, "PNG", PNG폭, int(PNG폭 * H / W))
            png.append(p)
        for k in range(1, d.Fonts.Count + 1):
            글꼴.add(d.Fonts.Item(k).Name)
        d.SaveAs(os.path.join(그림, "..", "발표_확인용.pdf"), 32)   # ppSaveAsPDF
    finally:
        d.Close()
    assert png, "PNG 가 0 장이다"
    return 오류, 권고, 글꼴, png


def 전체보기(png, 저장):
    """쪽 PNG 를 한 장에 모은다. 사람이 흐름과 들쭉날쭉한 배치를 한눈에 본다."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.image as mpimg
    열 = 4
    행 = (len(png) + 열 - 1) // 열
    fig, axs = plt.subplots(행, 열, figsize=(열 * 4, 행 * 2.5), squeeze=False)
    for k, ax in enumerate(axs.flat):
        if k >= len(png):
            ax.axis("off")
            continue
        ax.imshow(mpimg.imread(png[k]))
        ax.set_title("%d" % (k + 1), fontsize=9)
        ax.set_xticks([])
        ax.set_yticks([])
        for sp in ax.spines.values():        # white slides vanish on a white sheet
            sp.set_edgecolor("0.6")
    fig.tight_layout()
    fig.savefig(저장, dpi=110)
    plt.close(fig)


# ------------------------------------------------------------------ 본체

def main(pptx, 지정=None, 그림없이=False):
    assert os.path.isfile(pptx), "파일이 없다: %s" % pptx
    # PowerPoint COM resolves relative paths against its own cwd, not ours
    # (2026-09-22: Slide.Export failed with "찾을 수 없습니다" on a relative path)
    pptx = os.path.abspath(pptx)
    지정 = os.path.abspath(지정) if 지정 else None
    out = 출력폴더(pptx, 지정)
    그림 = os.path.join(out, "그림")
    os.makedirs(그림, exist_ok=True)

    쪽수, 오류, 권고, 글꼴 = 파일점검(pptx)
    png = []
    if not 그림없이:
        e, w, f, png = 그려점검(pptx, 그림)
        오류 += e
        권고 += w
        글꼴 |= f
        전체보기(png, os.path.join(그림, "00_전체보기.png"))

    있음 = 설치글꼴()
    없는글꼴 = sorted(x for x in 글꼴 if not 테마글꼴.match(x)
                  and x not in 있음 and 글꼴별칭.get(x) not in 있음)
    for x in 없는글꼴:
        오류.append("전체  이 PC 에 없는 글꼴 '%s'. 다른 글꼴로 바뀌어 줄바꿈이 달라진다" % x)

    줄 = ["# 발표자료 점검", "",
          "대상 `%s`" % pptx, "",
          "슬라이드 %d 장 · 오류 %d 건 · 권고 %d 건 · PNG %d 장" % (쪽수, len(오류), len(권고), len(png)),
          "", "쓰인 글꼴: %s" % (" · ".join(sorted(글꼴)) or "(테마 기본만)"), "",
          "## 오류", ""] + ["- " + x for x in 오류 or ["없음"]] + \
         ["", "## 권고", ""] + ["- " + x for x in 권고 or ["없음"]] + \
         ["", "## 사람이 볼 것", "",
          "`그림\\00_전체보기.png` 로 흐름을 본다. 제목만 이어 읽어 이야기가 되는지,",
          "쪽마다 배치가 들쭉날쭉하지 않은지, 숫자가 보고서와 같은지는 기계가 못 본다.", ""]
    open(os.path.join(out, "00_점검결과.md"), "w", encoding="utf-8").write("\n".join(줄))
    if png:
        목록 = ["# 그림 목록", "",
              "- `00_전체보기.png` 전체 흐름. 배치가 튀는 쪽, 글이 빽빽한 쪽을 찾는다"]
        목록 += ["- `%s` %d쪽. 넘침 · 겹침 · 잘림을 본다" % (os.path.basename(p), k)
               for k, p in enumerate(png, 1)]
        open(os.path.join(그림, "00_그림목록.md"), "w", encoding="utf-8").write("\n".join(목록) + "\n")

    print("슬라이드 %d 장 · PNG %d 장" % (쪽수, len(png)))
    for x in 오류:
        print("[오류]", x)
    for x in 권고:
        print("[권고]", x)
    print("결과 %s" % out)
    print("[결과] 오류 %d 건 · 권고 %d 건" % (len(오류), len(권고)))
    return 1 if 오류 else 0


if __name__ == "__main__":
    인자 = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not 인자:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(인자[0], 인자[1] if len(인자) > 1 else None,
                  그림없이="--그림없이" in sys.argv))
