r"""양식대조.py 가 글만 바꾼 판은 통과시키고, 틀을 건드린 판은 빠짐없이 잡는지 본다.

양식대조.py 를 고친 뒤 반드시 돌린다. 시험 파일은 임시 폴더에 만든다.

Rerun: C:\Users\SGJEONG99\anaconda3\envs\pyswmm\python.exe -X utf8 양식대조_시험.py
"""
import os
import subprocess
import sys
import tempfile

import docx
from docx.shared import Pt as DPt, Cm
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.join(tempfile.gettempdir(), "양식대조_시험")
TOOL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "양식대조.py")
os.makedirs(HERE, exist_ok=True)
P = lambda n: os.path.join(HERE, n)


def run(a, b):
    r = subprocess.run([sys.executable, "-X", "utf8", TOOL, a, b], capture_output=True,
                       text=True, encoding="utf-8")
    assert not r.stderr.strip(), r.stderr
    return r.returncode, r.stdout


def 채우기(tf, text):
    # replace text but keep run formatting (the right way)
    tf.paragraphs[0].runs[0].text = text


# ---------------------------------------------------------------- pptx original
Image.new("RGB", (120, 60), (20, 60, 120)).save(P("logo.png"))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for k in range(4):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "제목 %d" % (k + 1)
    b = s.placeholders[1]
    b.text = "본문 %d" % (k + 1)
    for r in b.text_frame.paragraphs[0].runs:
        r.font.name, r.font.size = "맑은 고딕", Pt(20)
    s.shapes.add_picture(P("logo.png"), Inches(11.8), Inches(0.2), Inches(1.2))
    ft = s.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(6), Inches(0.4))
    ft.text_frame.text = "기관명 | 대외비"
prs.save(P("orig.pptx"))

# good: text only
g = Presentation(P("orig.pptx"))
for k, s in enumerate(g.slides):
    채우기(s.shapes.title.text_frame, "고친 제목 %d" % (k + 1))
    채우기(s.placeholders[1].text_frame, "고친 본문 %d" % (k + 1))
g.save(P("good.pptx"))

# bad: layout edited, title moved, footer removed, fonts wiped
b = Presentation(P("orig.pptx"))
b.slide_layouts[1].placeholders[0].top += Inches(0.3)      # layout XML changed
s2, s3, s4 = b.slides[1], b.slides[2], b.slides[3]
s2.shapes.title.left += Inches(1)
foot = [sh for sh in s3.shapes if sh.has_text_frame and "대외비" in sh.text_frame.text][0]
foot._element.getparent().remove(foot._element)
s4.placeholders[1].text_frame.text = "본문 4"               # same text, formatting lost
b.save(P("bad.pptx"))

rc, out = run(P("orig.pptx"), P("good.pptx"))
print(out)
assert rc == 0 and "틀 오류 0 건" in out, "good pptx must pass"
assert "'글 바뀜': 8" in out, "text changes should be counted"

rc, out = run(P("orig.pptx"), P("bad.pptx"))
print(out)
expect = ["레이아웃 바뀜: ppt/slideLayouts/slideLayout2.xml",
          "2쪽  'Title 1' 위치 · 크기 바뀜",
          "3쪽  반복 요소가 빠지거나 바뀜",
          "4쪽  'Content Placeholder 2' 글꼴 바뀜"]
miss = [e for e in expect if e not in out]
assert rc == 1 and not miss, "missed: %s" % miss
# Editing the layout moves the inherited title on EVERY slide, including slide 1 that
# was never touched directly (measured 2026-09-22). The tool must show that cascade,
# but slide 1 must not get unrelated complaints (repeated elements, fonts)
assert "1쪽  'Title 1' 위치 · 크기 바뀜" in out, "layout cascade must be reported"
assert "1쪽  반복 요소" not in out and "1쪽  'Content Placeholder 2' 글꼴" not in out

# ---------------------------------------------------------------- docx
d = docx.Document()
d.sections[0].header.paragraphs[0].text = "기관명 과업 보고서"
d.add_heading("1. 서론", level=1)
d.add_paragraph("본문 첫 문단")
d.add_paragraph("본문 둘째 문단")
d.save(P("orig.docx"))

g = docx.Document(P("orig.docx"))
g.paragraphs[1].runs[0].text = "고친 첫 문단"
g.save(P("good.docx"))

b = docx.Document(P("orig.docx"))
b.sections[0].left_margin = Cm(1.5)
b.sections[0].header.paragraphs[0].text = "바뀐 머리글"
b.styles["Normal"].font.size = DPt(9)
b.paragraphs[2].runs[0].font.name = "Arial"
b.save(P("bad.docx"))

rc, out = run(P("orig.docx"), P("good.docx"))
print(out)
assert rc == 0 and "틀 오류 0 건" in out, "good docx must pass"

rc, out = run(P("orig.docx"), P("bad.docx"))
print(out)
expect = ["머리글 바뀜", "구역 1 쪽 설정 바뀜 (pgMar, 단위 twip): left 1800 -> 850", "스타일 정의 바뀜: Normal", "직접 준 서식 증가"]
miss = [e for e in expect if e not in out]
assert rc == 1 and not miss, "missed: %s" % miss

# ---------------------------------------------------------------- re-save must be silent
# Opening and saving in PowerPoint / Word rewrites XML (namespaces, rsid, localized
# style ids, dropped defaults). Before normalization this produced 13 and 118 false
# template errors (2026-09-22). Keep these two checks.
import win32com.client as win32
app = win32.Dispatch("PowerPoint.Application")
d = app.Presentations.Open(P("orig.pptx"), ReadOnly=True, WithWindow=False)
d.SaveCopyAs(P("resaved.pptx"))
d.Close()
rc, out = run(P("orig.pptx"), P("resaved.pptx"))
assert rc == 0, "PowerPoint re-save flagged:\n" + out

w = win32.Dispatch("Word.Application")
w.Visible = False
doc = w.Documents.Open(P("orig.docx"), ReadOnly=True)
doc.SaveAs2(P("resaved.docx"), FileFormat=16)            # wdFormatXMLDocument
doc.Close(False)
rc, out = run(P("orig.docx"), P("resaved.docx"))
assert rc == 0, "Word re-save flagged:\n" + out
print("OK: good versions pass, all 8 planted template changes caught, re-saves silent")
