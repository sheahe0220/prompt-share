r"""양식 대조. 고친 문서가 원본의 틀을 해치지 않았는지 원본과 맞대 본다.

왜 있나
  발표자료 · 보고서에는 쪽마다 같은 틀이 있다. 로고 · 머리글 · 바닥글 · 쪽 번호 ·
  제목 자리 · 글꼴 · 여백은 바뀌면 안 되고 글 내용만 바뀌어야 한다. 그런데 AI 는
  무엇이 고정인지 판단하지 못해 고칠 때마다 틀까지 건드린다(사용자 지적 2026-09-22).
  눈으로는 놓치므로 기계로 맞댄다. 글이 바뀐 것은 정상이고, 틀이 바뀐 것만 잡는다.

무엇을 보나
  pptx  오류  마스터 · 레이아웃 · 테마 · 슬라이드 크기 변경
             같은 레이아웃의 쪽마다 반복되던 요소(로고 · 바닥글 등)가 빠지거나 바뀜
             원래 있던 도형의 위치 · 크기 변경, 글꼴(이름 · 크기 · 굵기 · 색) 변경
        권고  도형 삭제 · 추가, 쪽 삭제 · 추가
  docx  오류  머리글 · 바닥글 · 테마 변경, 쪽 크기 · 여백 · 방향 변경, 스타일 정의 변경
        권고  스타일 추가 · 삭제, 글자에 직접 준 서식(글꼴 · 크기) 증가
  정상       글 내용 변경. 개수만 센다

실행
  C:\Users\SGJEONG99\anaconda3\envs\pyswmm\python.exe -X utf8 양식대조.py <원본> <수정본>
  종료코드 1 이면 틀이 바뀌었다. 결과는 수정본 옆 <이름>_양식대조.md 에도 남긴다.
시험
  양식대조_시험.py 가 글만 바꾼 판은 통과, 틀을 건드린 판은 잡는지 본다.
"""
import hashlib
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict

EMU_TOL = 12700          # 1 pt. rounding noise from re-saving


def 해시(z, n):
    """Hash of canonical XML. PowerPoint re-serializes parts on save (namespace
    declarations, attribute order) without changing meaning; raw bytes then differ
    although nothing changed (measured 2026-09-22: 12 layout parts flagged falsely)."""
    raw = z.read(n)
    if n.endswith((".xml", ".rels")):
        from lxml import etree
        t = etree.fromstring(raw, etree.XMLParser(remove_blank_text=True))
        for x in list(t.iter()):
            if x.tag == A + "endParaRPr":
                x.getparent().remove(x)
                continue
            for k in list(x.attrib):
                if etree.QName(k).localname in P소음속성:
                    del x.attrib[k]
            if x.tag == A + "fld":               # field results (date, slide no.) refresh on save
                for tt in x.findall(A + "t"):
                    tt.text = ""
        raw = etree.tostring(t, method="c14n")
    return hashlib.sha1(raw).hexdigest()


# PowerPoint save noise: spell-check state, language tags (measured 2026-09-22)
P소음속성 = {"smtClean", "smtId", "dirty", "err", "lang", "altLang", "noProof", "bmk"}


def 마스터테마(z):
    """Themes used by slide masters. Notes/handout masters have their own theme,
    which PowerPoint may regenerate on save; that is not the slide template."""
    out = set()
    for n in z.namelist():
        if re.match(r"ppt/slideMasters/_rels/slideMaster\d+\.xml\.rels$", n):
            for t in re.findall(r'Target="\.\./theme/(theme\d+\.xml)"', z.read(n).decode("utf-8")):
                out.add("ppt/theme/" + t)
    return out


def 부분대조(za, zb, 패턴):
    """Parts matching pattern: changed, added, removed."""
    a = {n for n in za.namelist() if re.match(패턴, n)}
    b = {n for n in zb.namelist() if re.match(패턴, n)}
    바뀜 = sorted(n for n in a & b if 해시(za, n) != 해시(zb, n))
    return 바뀜, sorted(b - a), sorted(a - b)


# ------------------------------------------------------------------ pptx

def 글꼴표(sh):
    """Run-level font properties of a shape, as a comparable tuple list."""
    from pptx.oxml.ns import qn
    out = []
    if not sh.has_text_frame:
        return out
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            f = r.font
            try:
                색 = str(f.color.rgb) if f.color and f.color.type is not None else None
            except AttributeError:
                색 = "theme"
            rPr = r._r.rPr
            ea = rPr.find(qn("a:ea")) if rPr is not None else None
            out.append((f.name, ea.get("typeface") if ea is not None else None,
                        "%.0fpt" % f.size.pt if f.size else None, f.bold, 색))
    return out


def 도형열쇠(sh):
    """Identity of a repeated element: kind + geometry + content (text or image hash)."""
    if sh.shape_type == 13:
        # SVG-only or linked pictures have no embedded blip; python-pptx raises
        # ValueError("no embedded image") (seen 2026-10-01 on an SVG graphic).
        try:
            내용 = "img:" + hashlib.sha1(sh.image.blob).hexdigest()[:12]
        except ValueError:
            내용 = "img-noblob:" + sh.name
    elif sh.has_text_frame:
        내용 = "txt:" + sh.text_frame.text.strip()
    else:
        내용 = "shape:%s" % sh.shape_type
    g = tuple(int(round((v or 0) / EMU_TOL)) for v in (sh.left, sh.top, sh.width, sh.height))
    return (내용, g)


def 설명(k):
    내용, g = k
    return "%s @ (%d, %d) pt" % (내용[:40], g[0], g[1])


def pptx대조(원본, 수정본):
    from pptx import Presentation
    오류, 권고, 정상 = [], [], Counter()
    za, zb = zipfile.ZipFile(원본), zipfile.ZipFile(수정본)
    슬라이드테마 = 마스터테마(za)
    for 이름, 패턴 in [("마스터", r"ppt/slideMasters/slideMaster\d+\.xml$"),
                     ("레이아웃", r"ppt/slideLayouts/slideLayout\d+\.xml$"),
                     ("테마", r"ppt/theme/theme\d+\.xml$")]:
        바뀜, 추가, 삭제 = 부분대조(za, zb, 패턴)
        for n in 바뀜:
            if 이름 == "테마" and n not in 슬라이드테마:
                권고.append("노트 · 유인물 테마 바뀜: %s (PowerPoint 가 저장하며 바꾸기도 한다)" % n)
            else:
                오류.append("%s 바뀜: %s" % (이름, n))
        for n in 추가 + 삭제:
            오류.append("%s %s: %s" % (이름, "추가" if n in 추가 else "삭제", n))

    A, B = Presentation(원본), Presentation(수정본)
    if (abs(A.slide_width - B.slide_width) > EMU_TOL or abs(A.slide_height - B.slide_height) > EMU_TOL):
        오류.append("슬라이드 크기 바뀜: %s -> %s" % ((A.slide_width, A.slide_height),
                                                  (B.slide_width, B.slide_height)))

    # Repeated elements per layout: on at least 80 % of that layout's slides. Not 100 %:
    # a cover or section slide without the logo would otherwise wipe out the set
    # (real 25-slide deck gave 0 repeated elements with the all-slides rule, 2026-09-22)
    쪽들 = defaultdict(list)
    for s in A.slides:
        쪽들[s.slide_layout.name].append({도형열쇠(sh) for sh in s.shapes})
    고정 = {}
    for lay, ks in 쪽들.items():
        if len(ks) < 2:
            continue
        c = Counter(k for kk in ks for k in kk)
        고정[lay] = {k for k, v in c.items() if v >= max(2, 0.8 * len(ks))}
    고정수 = sum(len(v) for v in 고정.values())
    원본열쇠 = {s.slide_id: {도형열쇠(sh) for sh in s.shapes} for s in A.slides}

    for i, s in enumerate(B.slides, 1):
        need = 고정.get(s.slide_layout.name)
        if not need:
            continue
        있음 = {도형열쇠(sh) for sh in s.shapes}
        if s.slide_id in 원본열쇠:              # same slide: what it had must stay
            for k in sorted((need & 원본열쇠[s.slide_id]) - 있음):
                오류.append("%d쪽  반복 요소가 빠지거나 바뀜 [%s]: %s" % (i, s.slide_layout.name, 설명(k)))
        else:                                    # new slide: should carry the usual ones
            for k in sorted(need - 있음):
                권고.append("%d쪽(새 쪽)  이 레이아웃에 늘 있던 요소가 없다: %s" % (i, 설명(k)))

    # same slide (by slide_id), same shape (by shape_id): geometry and fonts must hold
    a쪽 = {s.slide_id: s for s in A.slides}
    b번호 = {s.slide_id: k for k, s in enumerate(B.slides, 1)}
    for sid, sa in a쪽.items():
        if sid not in b번호:
            권고.append("쪽 삭제: 원본 %d쪽" % (list(a쪽).index(sid) + 1))
            continue
        i = b번호[sid]
        sb = B.slides[i - 1]
        if sb.slide_layout.name != sa.slide_layout.name:
            오류.append("%d쪽  레이아웃 바뀜: %s -> %s" % (i, sa.slide_layout.name, sb.slide_layout.name))
        ta = {sh.shape_id: sh for sh in sa.shapes}
        tb = {sh.shape_id: sh for sh in sb.shapes}
        for sid2, x in ta.items():
            if sid2 not in tb:
                권고.append("%d쪽  도형 삭제: '%s'" % (i, x.name))
                continue
            y = tb[sid2]
            ga = (x.left, x.top, x.width, x.height)
            gb = (y.left, y.top, y.width, y.height)
            if None not in ga + gb and any(abs(p - q) > EMU_TOL for p, q in zip(ga, gb)):
                오류.append("%d쪽  '%s' 위치 · 크기 바뀜: (%.0f, %.0f, %.0f, %.0f) -> (%.0f, %.0f, %.0f, %.0f) pt"
                          % ((i, x.name) + tuple(v / EMU_TOL for v in ga + gb)))
            # Only NEW font combinations count. A combination disappearing usually means
            # its text was deleted (real deck, 2026-09-22), which is a content edit
            fa, fb = 글꼴표(x), 글꼴표(y)
            새것 = set(fb) - set(fa)
            if fa and 새것:
                오류.append("%d쪽  '%s' 글꼴 바뀜: %s -> %s" % (i, x.name, sorted(set(fa) - set(fb), key=str)[:2],
                                                         sorted(새것, key=str)[:2]))
            if x.has_text_frame and y.has_text_frame and x.text_frame.text != y.text_frame.text:
                정상["글 바뀜"] += 1
        for sid2, y in tb.items():
            if sid2 not in ta:
                권고.append("%d쪽  도형 추가: '%s'" % (i, y.name))
    새쪽 = [k for sid, k in b번호.items() if sid not in a쪽]
    if 새쪽:
        권고.append("쪽 추가: %s쪽" % ", ".join(map(str, 새쪽)))
    요약 = "원본 %d쪽 · 수정본 %d쪽 · 반복 요소 %d 개(레이아웃 %d 종에서)" % (
        len(A.slides), len(B.slides), 고정수, len(고정))
    return 요약, 오류, 권고, 정상


# ------------------------------------------------------------------ docx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# Word re-saves with noise that means nothing for the template (measured 2026-09-22,
# Korean Word 2024): styleIds localized (Normal -> "a1", Header -> "a5"), rsid/paraId
# attributes added, whitespace and namespace declarations rewritten. Compare meaning only
소음속성 = re.compile(r"^(rsid.*|paraId|textId|Ignorable)$")
소음요소 = {W + t for t in ("proofErr", "lastRenderedPageBreak", "bookmarkStart", "bookmarkEnd",
                          "rsid", "rsids", "uiPriority", "semiHidden", "unhideWhenUsed",
                          "qFormat", "locked", "noProof", "lang")}
참조요소 = {W + t for t in ("pStyle", "rStyle", "tblStyle", "basedOn", "next", "numStyleLink",
                          "styleLink")}
# w:link names the paired "X Char" style, whose NAME Word localizes ("Heading 1 Char" ->
# "제목 1 Char"). It is bookkeeping, not formatting
소음요소 |= {W + "link"}
# Values equal to Word's defaults that Word drops on save (99 table styles, measured)
기본값 = {W + "tblInd": {"w": "0"},
       W + "tblCellMar": {"top": "0", "left": "108", "bottom": "0", "right": "108"}}


def 기본값인가(x):
    if x.tag == W + "tblInd":
        return x.get(W + "w") == "0"
    if x.tag == W + "tblCellMar":
        v = {etree_local(c.tag): c.get(W + "w") for c in x}
        return v == 기본값[x.tag]
    return False


def etree_local(tag):
    return tag.split("}")[-1]


def 스타일이름표(z):
    from lxml import etree
    t = etree.fromstring(z.read("word/styles.xml"))
    return {s.get(W + "styleId"): s.find(W + "name").get(W + "val")
            for s in t.iter(W + "style") if s.find(W + "name") is not None}


def 정규화(el, 이름표):
    """Canonical bytes of an element with Word's save noise removed and style
    references rewritten from localized ids to stable names."""
    from lxml import etree
    el = etree.fromstring(etree.tostring(el), etree.XMLParser(remove_blank_text=True))
    for x in list(el.iter()):
        if (x.tag in 소음요소 or 기본값인가(x)) and x is not el:
            x.getparent().remove(x)
            continue
        for k in list(x.attrib):
            if 소음속성.match(etree.QName(k).localname):
                del x.attrib[k]
        if x.tag in 참조요소 and x.get(W + "val") in 이름표:
            x.set(W + "val", 이름표[x.get(W + "val")])
    if W + "styleId" in el.attrib:
        del el.attrib[W + "styleId"]
    합치기(el)
    etree.cleanup_namespaces(el)
    return etree.tostring(el, method="c14n", exclusive=True)


def 합치기(el):
    """Merge neighbouring text runs with identical formatting. Word splits Korean
    text into one run per word on save (measured 2026-09-22); the look is the same."""
    from lxml import etree
    XS = "{http://www.w3.org/XML/1998/namespace}space"
    for p in el.iter(W + "p"):
        앞 = None
        for r in list(p):
            if r.tag != W + "r" or any(c.tag not in (W + "rPr", W + "t") for c in r):
                앞 = None
                continue
            서식 = r.find(W + "rPr")
            열쇠 = etree.tostring(서식, method="c14n") if 서식 is not None else b""
            if 앞 is not None and 앞[1] == 열쇠:
                t0, t1 = 앞[0].find(W + "t"), r.find(W + "t")
                if t0 is not None and t1 is not None:
                    t0.text = (t0.text or "") + (t1.text or "")
                    p.remove(r)
                    continue
            앞 = (r, 열쇠)
        for t in p.iter(W + "t"):
            t.attrib.pop(XS, None)


def 부분정규대조(za, zb, 패턴, 이름a, 이름b):
    from lxml import etree
    a = {n for n in za.namelist() if re.match(패턴, n)}
    b = {n for n in zb.namelist() if re.match(패턴, n)}
    바뀜 = sorted(n for n in a & b
                if 정규화(etree.fromstring(za.read(n)), 이름a) != 정규화(etree.fromstring(zb.read(n)), 이름b))
    return 바뀜 + sorted(a ^ b)


def 테마대조(za, zb):
    """Only colour and font schemes matter for the look."""
    from lxml import etree
    out = []
    for n in sorted(x for x in za.namelist() if re.match(r"word/theme/theme\d+\.xml$", x)):
        if n not in zb.namelist():
            out.append(n)
            continue
        ta, tb = etree.fromstring(za.read(n)), etree.fromstring(zb.read(n))
        for tag in ("clrScheme", "fontScheme"):
            ea, eb = ta.find(".//" + A + tag), tb.find(".//" + A + tag)
            if (ea is None) != (eb is None) or (ea is not None and 정규화(ea, {}) != 정규화(eb, {})):
                out.append("%s (%s)" % (n, tag))
    return out


def 스타일표(z):
    """Style definitions keyed by name (stable), not styleId (localized by Word)."""
    from lxml import etree
    이름표 = 스타일이름표(z)
    t = etree.fromstring(z.read("word/styles.xml"))
    # linked character styles ("X Char") are generated and renamed by Word; skip them
    return {s.find(W + "name").get(W + "val"): 정규화(s, 이름표)
            for s in t.iter(W + "style") if s.find(W + "name") is not None
            and not (s.get(W + "type") == "character" and s.find(W + "link") is not None)}


def 구역표(z):
    from lxml import etree
    t = etree.fromstring(z.read("word/document.xml"))
    out = []
    for sp in t.iter(W + "sectPr"):
        out.append({c.tag.replace(W, ""): {k.replace(W, ""): v for k, v in c.attrib.items()}
                    for c in sp if c.tag in (W + "pgSz", W + "pgMar", W + "cols", W + "docGrid")})
    return out


def 직접서식수(z):
    s = z.read("word/document.xml").decode("utf-8")
    return len(re.findall(r"<w:rFonts ", s)), len(re.findall(r"<w:sz ", s))


def docx대조(원본, 수정본):
    오류, 권고, 정상 = [], [], Counter()
    za, zb = zipfile.ZipFile(원본), zipfile.ZipFile(수정본)
    이름a, 이름b = 스타일이름표(za), 스타일이름표(zb)
    for 이름, 패턴 in [("머리글", r"word/header\d+\.xml$"), ("바닥글", r"word/footer\d+\.xml$")]:
        for n in 부분정규대조(za, zb, 패턴, 이름a, 이름b):
            오류.append("%s 바뀜: %s" % (이름, n))
    for n in 테마대조(za, zb):
        오류.append("테마 바뀜: %s" % n)
    sa, sb = 스타일표(za), 스타일표(zb)
    for k in sorted(set(sa) & set(sb)):
        if sa[k] != sb[k]:
            오류.append("스타일 정의 바뀜: %s" % k)
    for k in sorted(set(sb) - set(sa)):
        권고.append("스타일 추가: %s" % k)
    for k in sorted(set(sa) - set(sb)):
        권고.append("스타일 삭제: %s" % k)
    ca, cb = 구역표(za), 구역표(zb)
    if len(ca) != len(cb):
        오류.append("구역 수 바뀜: %d -> %d" % (len(ca), len(cb)))
    for k, (x, y) in enumerate(zip(ca, cb), 1):
        for tag in sorted(set(x) | set(y)):
            a, b = x.get(tag) or {}, y.get(tag) or {}
            차이 = ["%s %s -> %s" % (at, a.get(at), b.get(at))
                  for at in sorted(set(a) | set(b)) if a.get(at) != b.get(at)]
            if 차이:
                오류.append("구역 %d 쪽 설정 바뀜 (%s, 단위 twip): %s" % (k, tag, " · ".join(차이)))
    fa, fb = 직접서식수(za), 직접서식수(zb)
    if fb[0] > fa[0] or fb[1] > fa[1]:
        권고.append("글자에 직접 준 서식 증가: 글꼴 %d -> %d, 크기 %d -> %d. 스타일로 줘야 틀이 산다"
                  % (fa[0], fb[0], fa[1], fb[1]))
    import docx
    pa = [p.text for p in docx.Document(원본).paragraphs]
    pb = [p.text for p in docx.Document(수정본).paragraphs]
    정상["글 바뀐 문단"] = sum(1 for x, y in zip(pa, pb) if x != y) + abs(len(pa) - len(pb))
    요약 = "스타일 %d 개 · 구역 %d 개 · 문단 %d -> %d" % (len(sa), len(ca), len(pa), len(pb))
    return 요약, 오류, 권고, 정상


# ------------------------------------------------------------------ 본체

def main(원본, 수정본):
    for p in (원본, 수정본):
        assert os.path.isfile(p), "파일이 없다: %s" % p
    e1, e2 = os.path.splitext(원본)[1].lower(), os.path.splitext(수정본)[1].lower()
    assert e1 == e2, "형식이 다르다: %s / %s" % (e1, e2)
    if e1 in (".pptx", ".potx"):
        요약, 오류, 권고, 정상 = pptx대조(원본, 수정본)
    elif e1 == ".docx":
        요약, 오류, 권고, 정상 = docx대조(원본, 수정본)
    else:
        raise SystemExit("pptx · docx 만 본다. hwpx 는 14_HWPX.md 의 방식으로 확인한다: %s" % e1)

    줄 = ["# 양식 대조", "", "원본 `%s`" % 원본, "수정본 `%s`" % 수정본, "", 요약,
          "틀 오류 %d 건 · 권고 %d 건 · 정상 변경 %s" % (len(오류), len(권고), dict(정상) or "없음"),
          "", "## 틀 오류 (고친다)", ""] + ["- " + x for x in 오류 or ["없음"]] + \
         ["", "## 권고 (의도한 것인지 본다)", ""] + ["- " + x for x in 권고 or ["없음"]] + [""]
    결과 = os.path.splitext(수정본)[0] + "_양식대조.md"
    open(결과, "w", encoding="utf-8").write("\n".join(줄))
    print(요약)
    for x in 오류:
        print("[틀]", x)
    for x in 권고:
        print("[권고]", x)
    print("정상 변경", dict(정상) or "없음")
    print("결과", 결과)
    print("[결과] 틀 오류 %d 건 · 권고 %d 건" % (len(오류), len(권고)))
    return 1 if 오류 else 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2]))
